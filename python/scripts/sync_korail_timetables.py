from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW_ROOT = ROOT / "data" / "raw" / "timetables"
PROCESSED_ROOT = ROOT / "data" / "processed" / "timetables"
BOARD_URL = "https://www.korail.com/com/userBoard.do?mode=list&schBcid=ticketTable"
FILE_BASE = "https://www.korail.com/file/cubedata/COMMON/"
USER_AGENT = "SunSeat-timetable-sync/1.0"

TARGETS = {
    "ktx": re.compile(r"KTX\s*시간표", re.IGNORECASE),
    "general": re.compile(r"일반열차\s*시간표", re.IGNORECASE),
    "itx-cheongchun": re.compile(r"ITX[-\s]*청춘.*(?:시간표|시각표)", re.IGNORECASE),
}


def safe_name(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", value).strip("_")[:120]


def board_items(payload: object) -> list[dict]:
    if not isinstance(payload, dict):
        raise ValueError("Korail board response is not an object")
    items = payload.get("boardList")
    if not isinstance(items, list):
        raise ValueError("Korail board response does not contain boardList")
    return [item for item in items if isinstance(item, dict)]


def choose_latest(items: list[dict], pattern: re.Pattern[str]) -> dict | None:
    candidates = [item for item in items if pattern.search(str(item.get("bdTitle") or "")) and item.get("fileId")]
    candidates.sort(key=lambda item: str(item.get("regdt") or ""), reverse=True)
    return candidates[0] if candidates else None


def inspect_workbook(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(min_row=1, max_row=min(30, worksheet.max_row), values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row[:30]]
            if any(values):
                rows.append(values)
        sheets.append({"title": worksheet.title, "maxRow": worksheet.max_row, "maxColumn": worksheet.max_column, "sampleRows": rows[:12]})
    workbook.close()
    return {"sheetNames": [sheet["title"] for sheet in sheets], "sheets": sheets}


def main() -> None:
    generated_at = datetime.now(timezone.utc)
    response = requests.get(BOARD_URL, headers={"User-Agent": USER_AGENT, "Accept": "application/json"}, timeout=60)
    response.raise_for_status()
    payload = response.json()
    items = board_items(payload)
    version = generated_at.strftime("v%Y%m%d%H%M%S")
    raw_version = RAW_ROOT / version
    processed_version = PROCESSED_ROOT / version
    raw_version.mkdir(parents=True, exist_ok=True)
    processed_version.mkdir(parents=True, exist_ok=True)
    records = []
    for key, pattern in TARGETS.items():
        item = choose_latest(items, pattern)
        if item is None:
            records.append({"key": key, "status": "NOT_FOUND"})
            continue
        file_ids = item.get("fileId") if isinstance(item.get("fileId"), list) else [item.get("fileId")]
        file_id = next((str(value) for value in file_ids if value), None)
        if not file_id:
            records.append({"key": key, "status": "NO_FILE_ID", "bdIdx": item.get("bdIdx"), "bdTitle": item.get("bdTitle")})
            continue
        download_url = urljoin(FILE_BASE, file_id)
        file_response = requests.get(download_url, headers={"User-Agent": USER_AGENT}, timeout=120)
        file_response.raise_for_status()
        filename = f"{key}-{safe_name(str(item.get('bdTitle') or item.get('bdIdx')))}.xlsx"
        raw_path = raw_version / filename
        raw_path.write_bytes(file_response.content)
        inspection = inspect_workbook(raw_path)
        records.append({
            "key": key,
            "status": "DOWNLOADED",
            "bdIdx": item.get("bdIdx"),
            "bdTitle": item.get("bdTitle"),
            "regdt": item.get("regdt"),
            "fileId": file_id,
            "sourceUrl": download_url,
            "rawFile": str(raw_path.relative_to(ROOT)).replace("\\", "/"),
            "inspection": inspection,
        })
    manifest = {"provider": "KORAIL", "sourceUrl": BOARD_URL, "generatedAt": generated_at.isoformat(), "version": version, "records": records}
    (processed_version / "timetable-index.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (PROCESSED_ROOT / "active.json").write_text(json.dumps({"version": version, "index": f"{version}/timetable-index.json", "activatedAt": generated_at.isoformat()}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"timetable sync failed: {error}", file=sys.stderr)
        raise
