from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
TIMETABLE_ROOT = ROOT / "data" / "processed" / "timetables"


def text(value: object) -> str:
    return str(value or "").replace("\n", " ").strip()


def station_text(value: object) -> str:
    return re.sub(r"\s+", "", text(value))


def train_number(value: object) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return str(int(value))
    value = text(value)
    return value if re.fullmatch(r"[A-Za-z]?\d{1,5}", value) else None


def format_time(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    value = text(value)
    if not value or value in {"00:00", "00:00:00", "-"}:
        return None
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", value)
    if not match:
        return None
    return f"{int(match.group(1)):02d}:{match.group(2)}"


def group_headers(rows: list[list[object]]) -> list[tuple[int, int]]:
    headers = []
    for row_index, row in enumerate(rows):
        positions = [index for index, value in enumerate(row) if text(value) == "열차번호"]
        headers.extend((row_index, index) for index in positions)
    return headers


def direction_for(rows: list[list[object]], row_index: int, column: int) -> str:
    for candidate_row in rows[max(0, row_index - 3):row_index + 1]:
        for value in candidate_row[max(0, column - 2):column + 5]:
            value_text = text(value)
            if "하행" in value_text or "From" in value_text and "to" in value_text:
                return "DOWN"
            if "상행" in value_text:
                return "UP"
    return "UNKNOWN"


def parse_sheet(rows: list[list[object]], sheet_name: str) -> list[dict]:
    output = []
    headers = group_headers(rows)
    for header_position, (header_row, start_column) in enumerate(headers):
        next_group = headers[header_position + 1][1] if header_position + 1 < len(headers) else len(rows[header_row])
        if next_group <= start_column + 2:
            continue
        stations: list[tuple[int, str]] = []
        for column in range(start_column + 2, next_group):
            name = station_text(rows[header_row][column] if column < len(rows[header_row]) else "")
            if name and name not in {"비고", "Remark", "Remarks"}:
                stations.append((column, name))
        if not stations:
            continue
        direction = direction_for(rows, header_row, start_column)
        for row_index in range(header_row + 1, len(rows)):
            row = rows[row_index]
            number = train_number(row[start_column] if start_column < len(row) else None)
            if number is None:
                continue
            grade = text(row[start_column + 1] if start_column + 1 < len(row) else "") or None
            stops = []
            for column, station in stations:
                value = format_time(row[column] if column < len(row) else None)
                if value:
                    stops.append({"station": station, "time": value})
            if len(stops) < 2:
                continue
            output.append({
                "sheet": sheet_name,
                "direction": direction,
                "trainNumber": number,
                "trainGradeName": grade,
                "origin": stops[0]["station"],
                "destination": stops[-1]["station"],
                "stops": stops,
            })
    return output


def parse_transposed_sheet(rows: list[list[object]], sheet_name: str) -> list[dict]:
    number_row_index = next((index for index, row in enumerate(rows) if text(row[0] if row else "") == "열차번호"), None)
    if number_row_index is None:
        return []
    number_row = rows[number_row_index]
    columns = [(column, train_number(value)) for column, value in enumerate(number_row[1:], start=1)]
    columns = [(column, number) for column, number in columns if number]
    if len(columns) < 2:
        return []
    type_row = next((row for row in rows[max(0, number_row_index - 3):number_row_index] if text(row[0] if row else "") == "열차종별"), None)
    default_grade = "ITX-청춘" if "청춘" in sheet_name else None
    grades = {column: text(type_row[column]) if type_row and column < len(type_row) else default_grade for column, _ in columns}
    direction = "UP" if "상행" in sheet_name else "DOWN" if "하행" in sheet_name else "UNKNOWN"
    stops_by_column = {column: [] for column, _ in columns}
    current_station = None
    for row in rows[number_row_index + 1:]:
        first = station_text(row[0] if row else "")
        if first and first not in {"열차종별", "열차번호", "시발역", "종착역", "비고", "Remark", "Remarks"}:
            current_station = first
        if not current_station:
            continue
        for column, _ in columns:
            value = format_time(row[column] if column < len(row) else None)
            if value:
                stops_by_column[column].append({"station": current_station, "time": value})
    output = []
    for column, number in columns:
        stops = stops_by_column[column]
        if len(stops) < 2:
            continue
        output.append({
            "sheet": sheet_name,
            "direction": direction,
            "trainNumber": number,
            "trainGradeName": grades.get(column) or default_grade,
            "origin": stops[0]["station"],
            "destination": stops[-1]["station"],
            "stops": stops,
        })
    return output


def parse_workbook(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    output = []
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        parsed = parse_sheet(rows, worksheet.title)
        if not parsed:
            parsed = parse_transposed_sheet(rows, worksheet.title)
        output.extend(parsed)
    workbook.close()
    return output


def main() -> None:
    active = json.loads((TIMETABLE_ROOT / "active.json").read_text(encoding="utf-8"))
    version = active["version"]
    index_path = TIMETABLE_ROOT / version / "timetable-index.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    records = []
    all_trains = []
    for record in index.get("records", []):
        if record.get("status") != "DOWNLOADED":
            records.append(record)
            continue
        source = ROOT / record["rawFile"]
        trains = parse_workbook(source)
        parsed_path = TIMETABLE_ROOT / version / f"{record['key']}-trains.json"
        parsed_path.write_text(json.dumps({"source": record, "trains": trains}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        records.append({**record, "parsedFile": str(parsed_path.relative_to(ROOT)).replace("\\", "/"), "parsedTrainCount": len(trains)})
        all_trains.extend({**train, "timetableKey": record["key"], "sourceFile": record["rawFile"]} for train in trains)
    result = {"provider": "KORAIL", "version": version, "generatedAt": datetime.now().astimezone().isoformat(), "records": records, "trainCount": len(all_trains), "trains": all_trains}
    (TIMETABLE_ROOT / version / "parsed-timetables.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"version": version, "trainCount": len(all_trains), "records": [{"key": item.get("key"), "status": item.get("status"), "parsedTrainCount": item.get("parsedTrainCount", 0)} for item in records]}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"timetable parse failed: {error}", file=sys.stderr)
        raise
